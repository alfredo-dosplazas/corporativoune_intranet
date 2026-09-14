from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import or_

from apps.sae.db import sae_session
from apps.sae.models_sae import get_sae_models


class Command(BaseCommand):
    help = "Exporta compras de un proveedor a Excel en la carpeta Downloads"

    def add_arguments(self, parser):
        parser.add_argument(
            '--empresa',
            type=str,
            default='abraham',
            help='Clave de la empresa (ej: abraham, dominum)'
        )
        parser.add_argument(
            '--proveedor',
            type=str,
            required=True,
            help='Clave exacta o nombre del proveedor (LIKE)'
        )
        parser.add_argument(
            '--fecha-inicio',
            type=str,
            required=True,
            help='Fecha inicial en formato YYYY-MM-DD'
        )
        parser.add_argument(
            '--fecha-fin',
            type=str,
            required=True,
            help='Fecha final en formato YYYY-MM-DD'
        )

    def handle(self, *args, **options):
        empresa_key = options['empresa']
        busqueda_prov = options['proveedor'].strip()

        try:
            f_inicio = datetime.strptime(options['fecha_inicio'], '%Y-%m-%d').date()
            f_fin = datetime.strptime(options['fecha_fin'], '%Y-%m-%d').date()
        except ValueError:
            self.stderr.write(self.style.ERROR("Formato de fecha inválido. Usa YYYY-MM-DD."))
            return

        with sae_session(empresa_key) as (session, suffix):
            models = get_sae_models(suffix)

            filtro_proveedor = or_(
                models.Compra.clave_proveedor == busqueda_prov,
                models.Proveedor.nombre.ilike(f"%{busqueda_prov}%")
            )

            query = (
                session.query(models.PartidaCompra)
                .join(models.PartidaCompra.compra)
                .outerjoin(models.PartidaCompra.producto)
                .outerjoin(models.Compra.proveedor)
                .filter(
                    filtro_proveedor,
                    models.Compra.fecha >= f_inicio,
                    models.Compra.fecha <= f_fin,
                    models.Compra.status != 'C'
                )
                .order_by(models.Compra.fecha.asc(), models.Compra.folio.asc())
            )

            registros = query.all()

            if not registros:
                self.stdout.write(
                    self.style.WARNING(f"No se encontraron compras para '{busqueda_prov}' en el periodo especificado.")
                )
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Compras"

            # Encabezados exactos a la imagen
            headers = [
                "DOCUMENTO", "CLAVE_PROVEEDOR", "PROVEEDOR", "PREFJIO",
                "ALMACEN", "CLAVE_PRODUCTO", "PRODUCTO", "FECHA",
                "CANTIDAD", "COSTO", "TOTAL_COSTO"
            ]
            ws.append(headers)

            # Estilo del Encabezado
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Llenado de filas
            for reg in registros:
                compra = reg.compra
                prod = reg.producto
                prov = compra.proveedor

                subtotal_calc = (reg.cantidad or 0.0) * (reg.costo or 0.0)
                total_costo = getattr(reg, 'importe', None) or subtotal_calc

                ws.append([
                    compra.folio,
                    compra.clave_proveedor.strip() if compra.clave_proveedor else '',
                    prov.nombre.strip() if prov and prov.nombre else '',
                    "AG",  # Prefijo estático
                    compra.num_alma or 1,
                    reg.cve_art.strip() if reg.cve_art else '',
                    prod.descripcion.strip() if prod and prod.descripcion else '',
                    compra.fecha.strftime('%d/%m/%Y') if compra.fecha else '',
                    reg.cantidad,
                    reg.costo,
                    total_costo
                ])

            # Formato de celdas
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                row[8].number_format = '#,##0'       # CANTIDAD
                row[9].number_format = '#,##0.00'    # COSTO
                row[10].number_format = '#,##0.00'   # TOTAL_COSTO

            # Ancho dinámico de columnas
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            downloads_dir = Path.home() / "Downloads"
            downloads_dir.mkdir(parents=True, exist_ok=True)

            nombre_limpio = "".join(c for c in busqueda_prov if c.isalnum() or c in (' ', '_', '-')).strip()
            filename = f"Compras_{empresa_key}_{nombre_limpio}_{f_inicio}_al_{f_fin}.xlsx"
            filepath = downloads_dir / filename

            wb.save(filepath)

            self.stdout.write(
                self.style.SUCCESS(f"Reporte generado exitosamente en: {filepath}")
            )