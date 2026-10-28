import os
import openpyxl
from openpyxl.drawing.image import Image as OpenPyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
from openpyxl.utils import get_column_letter
from django.utils.timezone import now


def build_inventario_excel(data, active_fields, campos_mapa, logo_path=None):
    wb = openpyxl.Workbook()

    wb.properties.title = "Reporte de Inventario - Encorto"
    wb.properties.subject = "Existencias e Inventario"
    wb.properties.creator = "Sistema Encorto"

    ws = wb.active
    ws.title = "Inventario"
    ws.views.sheetView[0].showGridLines = True

    # ---------------------------------------------------------
    # Estilos inspirados en el tema Encorto
    # ---------------------------------------------------------
    font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")  # Secondary Encorto
    font_subtitle = Font(name="Calibri", size=9, italic=True, color="4B5563")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="1F2937")
    font_total = Font(name="Calibri", size=10, bold=True, color="1E293B")

    fill_header = PatternFill(start_color="28A71A", end_color="28A71A", fill_type="solid")  # Primary Encorto
    fill_total = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    border_thin = Side(border_style="thin", color="D1D5DB")
    border_thick_top = Side(border_style="thin", color="1E293B")
    border_double_bottom = Side(border_style="double", color="1E293B")

    border_data = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    border_totals = Border(top=border_thick_top, bottom=border_double_bottom, left=border_thin, right=border_thin)

    # ---------------------------------------------------------
    # 1. Encabezado principal y Logo
    # ---------------------------------------------------------
    has_logo = False
    if logo_path and os.path.exists(logo_path):
        try:
            img = OpenPyxlImage(logo_path)
            # Redimensionar logo para que encaje bien en las primeras filas
            img.width = 110
            img.height = 35
            ws.add_image(img, "A1")
            has_logo = True
        except Exception:
            pass

    if has_logo:
        # Si hay logo, colocamos el título un poco desplazado o en la columna C
        ws["C1"] = "REPORTE DE INVENTARIO - ENCORTO"
        ws["C1"].font = font_title
        ws["C2"] = f"Fecha de generación: {now().strftime('%d/%m/%Y %H:%M')}"
        ws["C2"].font = font_subtitle
        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 18
    else:
        ws["A1"] = "REPORTE DE INVENTARIO - ENCORTO"
        ws["A1"].font = font_title
        ws["A2"] = f"Fecha de generación: {now().strftime('%d/%m/%Y %H:%M')}"
        ws["A2"].font = font_subtitle

    ws.append([])  # Fila 3 vacía

    # ---------------------------------------------------------
    # 2. Encabezados de tabla
    # ---------------------------------------------------------
    headers = [campos_mapa[field] for field in active_fields if field in campos_mapa]
    ws.append(headers)

    header_row = 4
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # ---------------------------------------------------------
    # 3. Llenado de datos
    # ---------------------------------------------------------
    start_data_row = 5
    for item in data:
        row_values = [item.get(field, '') for field in active_fields]
        ws.append(row_values)

    end_data_row = start_data_row + len(data) - 1

    # ---------------------------------------------------------
    # 4. Formateo de Celdas
    # ---------------------------------------------------------
    for r_idx in range(start_data_row, end_data_row + 1):
        for c_idx, field in enumerate(active_fields, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_data
            cell.border = border_data

            if 'precio' in field or 'costo' in field or 'mayoreo' in field:
                cell.alignment = align_right
                cell.number_format = '"$"#,##0.00'
            elif field in ['existencia', 'cantidad']:
                cell.alignment = align_right
                cell.number_format = '#,##0.00'
            elif field == 'articulo_descripcion':
                cell.alignment = align_left
            else:
                cell.alignment = align_center

    # ---------------------------------------------------------
    # 5. Fila de Totales
    # ---------------------------------------------------------
    if data:
        total_row_idx = end_data_row + 1
        ws.cell(row=total_row_idx, column=1, value="TOTALES").font = font_total

        for c_idx, field in enumerate(active_fields, start=1):
            cell = ws.cell(row=total_row_idx, column=c_idx)
            cell.font = font_total
            cell.fill = fill_total
            cell.border = border_totals

            col_letter = get_column_letter(c_idx)

            if field in ['existencia', 'costo_total']:
                cell.value = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                cell.alignment = align_right
                if field == 'existencia':
                    cell.number_format = '#,##0.00'
                else:
                    cell.number_format = '"$"#,##0.00'

    ws.freeze_panes = "A5"

    # ---------------------------------------------------------
    # 6. Anchos predefinidos por columna
    # ---------------------------------------------------------
    anchos_default = {
        'almacen_nombre': 18,
        'articulo_clave': 14,
        'status_producto': 10,
        'articulo_descripcion': 45,
        'linea': 12,
        'ctrl_alm': 15,
        'existencia': 15,
        'precio_lista': 16,
        'mayoreo': 16,
        'costo_promedio': 16,
        'costo_total': 18,
    }

    for col_idx, field in enumerate(active_fields, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = anchos_default.get(field, 16)

    return wb