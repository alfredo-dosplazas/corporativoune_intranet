import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.contrib import messages

from django.db import transaction
from django.db.models import Count
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from .models import DetalleVoto, EventoVotacion, OpcionParticipante, VotoSesion
from ..core.decorators import modulo_required


@modulo_required('Votaciones')
def votar_evento(request, evento_id):
    evento = get_object_or_404(EventoVotacion, id=evento_id, activo=True)
    ahora = timezone.now()
    user = request.user

    # -------------------------------------------------------------
    # 1. CONTROL DE TIEMPOS DE VOTACIÓN
    # -------------------------------------------------------------
    # Excepción: Los superusuarios pueden saltarse el bloqueo para pruebas
    if not user.is_superuser:
        # A) Si la votación aún no ha comenzado
        if ahora < evento.fecha_inicio:
            return render(
                request,
                'apps/votaciones/proximamente.html',
                {'evento': evento},
            )

        # B) Si la votación ya finalizó -> Redirigir a resultados
        if ahora > evento.fecha_fin:
            messages.info(
                request,
                'El tiempo para votar en este evento ha finalizado.',
            )
            return redirect('votaciones:resultados_evento', evento_id=evento.id)

    # -------------------------------------------------------------
    # 2. PROCESAMIENTO DEL VOTO (POST)
    # -------------------------------------------------------------
    criterios = evento.criterios.all()
    participantes = evento.participantes.all()

    if request.method == 'POST':
        identificador = request.POST.get('identificador', '').strip()

        if evento.requiere_identificador and not identificador:
            messages.error(
                request,
                'Debes ingresar tu número de empleado o ID para votar.',
            )
            return render(
                request,
                'apps/votaciones/votar.html',
                {
                    'evento': evento,
                    'criterios': criterios,
                    'participantes': participantes,
                },
            )

        if (
            evento.requiere_identificador
            and VotoSesion.objects.filter(
                evento=evento, identificador=identificador
            ).exists()
        ):
            messages.error(
                request,
                f'El identificador "{identificador}" ya ha registrado su voto en este concurso.',
            )
            return render(
                request,
                'apps/votaciones/votar.html',
                {
                    'evento': evento,
                    'criterios': criterios,
                    'participantes': participantes,
                },
            )

        ip = request.META.get('HTTP_X_FORWARDED_FOR')
        if ip:
            ip = ip.split(',')[0]
        else:
            ip = request.META.get('REMOTE_ADDR')

        try:
            with transaction.atomic():
                sesion_voto = VotoSesion.objects.create(
                    evento=evento,
                    identificador=identificador
                    or f'ANON-{VotoSesion.objects.count() + 1}',
                    ip_address=ip,
                )

                for c in criterios:
                    key_input = f'categoria_{c.id}'
                    participante_id = request.POST.get(key_input)

                    if participante_id:
                        participante_obj = participantes.filter(
                            id=participante_id
                        ).first()
                        if participante_obj:
                            DetalleVoto.objects.create(
                                sesion_voto=sesion_voto,
                                participante=participante_obj,
                                criterio=c,
                                valor=1,
                            )

            messages.success(
                request,
                '¡Tu voto ha sido registrado con éxito! Gracias por participar.',
            )
            return redirect(f'{request.path}?exito=1')

        except Exception as e:
            messages.error(
                request,
                f'Ocurrió un error al procesar tu voto. Inténtalo de nuevo. Error: {str(e)}',
            )

    exito = request.GET.get('exito') == '1'

    return render(
        request,
        'apps/votaciones/votar.html',
        {
            'evento': evento,
            'criterios': criterios,
            'participantes': participantes,
            'exito': exito,
        },
    )


@modulo_required('Votaciones')
def resultados_evento(request, evento_id):
    evento = get_object_or_404(EventoVotacion, id=evento_id)
    criterios = evento.criterios.all()
    participantes = evento.participantes.all()
    total_votos = evento.votosesion_set.count()

    resultados_por_criterio = []
    ganadores_asignados = set()

    for criterio in criterios:
        votos_criterio = (
            DetalleVoto.objects.filter(
                sesion_voto__evento=evento, criterio=criterio
            )
            .values('participante')
            .annotate(total_votos=Count('id'))
            .order_by('-total_votos')
        )

        ranking = []
        for v in votos_criterio:
            part = OpcionParticipante.objects.get(id=v['participante'])
            ranking.append({'participante': part, 'votos': v['total_votos']})

        ganador = None
        for r in ranking:
            if r['participante'] not in ganadores_asignados:
                ganador = r
                ganadores_asignados.add(r['participante'])
                break

        if not ganador and ranking:
            ganador = ranking[0]

        resultados_por_criterio.append(
            {'criterio': criterio, 'ganador': ganador, 'ranking': ranking}
        )

    sesiones_votos = evento.votosesion_set.prefetch_related(
        'detalles__criterio', 'detalles__participante'
    ).order_by('-fecha_emision')

    votos_detalle = []
    for sesion in sesiones_votos:
        votos_usuario = {}
        for detalle in sesion.detalles.all():
            votos_usuario[detalle.criterio_id] = detalle.participante.nombre

        votos_detalle.append(
            {
                'identificador': sesion.identificador,
                'fecha': sesion.fecha_emision,
                'ip': sesion.ip_address,
                'elecciones': votos_usuario,
            }
        )

    acumulado_participantes = []
    for part in participantes:
        conteo_criterios = {}
        total_acumulado = 0
        for c in criterios:
            cant_votos = DetalleVoto.objects.filter(
                sesion_voto__evento=evento, criterio=c, participante=part
            ).count()
            conteo_criterios[c.id] = cant_votos
            total_acumulado += cant_votos

        acumulado_participantes.append(
            {
                'participante': part,
                'conteo': conteo_criterios,
                'total_acumulado': total_acumulado,
            }
        )

    acumulado_participantes.sort(
        key=lambda x: x['total_acumulado'], reverse=True
    )

    return render(
        request,
        'apps/votaciones/resultados.html',
        {
            'evento': evento,
            'criterios': criterios,
            'resultados': resultados_por_criterio,
            'votos_detalle': votos_detalle,
            'acumulado_participantes': acumulado_participantes,
            'total_votos': total_votos,
        },
    )


@modulo_required('Votaciones')
def exportar_reporte_excel(request, evento_id):
    evento = get_object_or_404(EventoVotacion, id=evento_id)
    criterios = list(evento.criterios.all())
    participantes = list(evento.participantes.all())
    sesiones_votos = evento.votosesion_set.prefetch_related(
        'detalles__criterio', 'detalles__participante'
    ).order_by('-fecha_emision')

    wb = openpyxl.Workbook()

    # Estilos compartidos
    fill_header = PatternFill(
        start_color="311B92", end_color="311B92", fill_type="solid"
    )
    fill_accent = PatternFill(
        start_color="E8EAF6", end_color="E8EAF6", fill_type="solid"
    )
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_title = Font(name="Calibri", size=14, bold=True, color="311B92")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # -------------------------------------------------------------
    # HOJA 1: RESUMEN ACUMULADO Y TOTALES
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Resumen de Votaciones"

    ws1.append([f"REPORTE GENERAL DE VOTACIÓN - {evento.titulo.upper()}"])
    ws1.cell(row=1, column=1).font = font_title
    ws1.append([f"Fecha de generación: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws1.append([])

    headers_h1 = ["Participante / Disfraz"] + [c.texto for c in criterios] + ["TOTAL VOTOS"]
    ws1.append(headers_h1)

    row_header_idx = 4
    for col_num in range(1, len(headers_h1) + 1):
        cell = ws1.cell(row=row_header_idx, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for part in participantes:
        row = [part.nombre]
        total_part = 0
        for c in criterios:
            cnt = DetalleVoto.objects.filter(
                sesion_voto__evento=evento, criterio=c, participante=part
            ).count()
            row.append(cnt)
            total_part += cnt
        row.append(total_part)
        ws1.append(row)

        current_row = ws1.max_row
        for col_num in range(1, len(row) + 1):
            cell = ws1.cell(row=current_row, column=col_num)
            cell.border = thin_border
            if col_num == 1:
                cell.font = font_bold
            else:
                cell.alignment = Alignment(horizontal="center")
            if col_num == len(row):
                cell.fill = fill_accent
                cell.font = font_bold

    # -------------------------------------------------------------
    # HOJA 2: AUDITORÍA INDIVIDUAL (QUIÉN VOTÓ POR QUIÉN)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Auditoría de Votos")

    ws2.append([f"BITÁCORA DE AUDITORÍA DETALLADA - {evento.titulo.upper()}"])
    ws2.cell(row=1, column=1).font = font_title
    ws2.append([])

    headers_h2 = ["Votante (ID / Nombre)", "Fecha", "Hora", "Dirección IP"] + [f"Voto: {c.texto}" for c in criterios]
    ws2.append(headers_h2)

    row_header2_idx = 3
    for col_num in range(1, len(headers_h2) + 1):
        cell = ws2.cell(row=row_header2_idx, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for sesion in sesiones_votos:
        elecciones = {d.criterio_id: d.participante.nombre for d in sesion.detalles.all()}
        row = [
            sesion.identificador,
            sesion.fecha_emision.strftime("%Y-%m-%d"),
            sesion.fecha_emision.strftime("%H:%M:%S"),
            sesion.ip_address or "N/A",
        ]
        for c in criterios:
            row.append(elecciones.get(c.id, "Sin Selección"))

        ws2.append(row)
        current_row = ws2.max_row
        for col_num in range(1, len(row) + 1):
            cell = ws2.cell(row=current_row, column=col_num)
            cell.border = thin_border
            if col_num in [2, 3, 4]:
                cell.alignment = Alignment(horizontal="center")

    # Ajuste automático del ancho de columnas en ambas hojas
    for sheet in [ws1, ws2]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Respuesta de descarga HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"reporte_votacion_evento_{evento_id}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)

    return response